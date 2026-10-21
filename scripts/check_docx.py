"""精确检查：仅在非等宽字体（正文/表格正文）的 run 中查找 Markdown 残留标记。

代码块与行内代码使用 Consolas，其中的 ** / | / ``` 属于正常内容，不计为残留。
"""
import glob
import os

from docx import Document

OUT = r'C:\Users\wangchuan\Desktop\软件工程实践实验'
MARKERS = ['**', '```', '](', '## ', '### ']


def run_font_name(run):
    try:
        return (run.font.name or '') or ''
    except Exception:
        return ''


def scan_paragraph(paragraph):
    hits = []
    if run_font_name(paragraph.runs[0]) if paragraph.runs else False:
        pass
    for run in paragraph.runs:
        if 'Consolas' in run_font_name(run):
            continue
        for marker in MARKERS:
            if marker in run.text:
                hits.append((marker, run.text.strip()[:60]))
    return hits


total_bad = 0
for path in sorted(glob.glob(os.path.join(OUT, '*.docx'))):
    if '任务划分' in path:
        continue
    doc = Document(path)
    bad = []
    for paragraph in doc.paragraphs:
        bad.extend(scan_paragraph(paragraph))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    bad.extend(scan_paragraph(paragraph))
    total_bad += len(bad)
    size_kb = round(os.path.getsize(path) / 1024, 1)
    print('%-38s 段落=%-4d 表格=%-3d %6.1fKB  非代码残留=%d'
          % (os.path.basename(path), len(doc.paragraphs), len(doc.tables), size_kb, len(bad)))
    for marker, text in bad[:3]:
        print('    !!', repr(marker), text)

print('总残留：', total_bad)
