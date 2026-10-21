#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Markdown -> Word(.docx) 课程文档转换器。

用法：
    python md2docx.py <源目录> <输出目录> [--toc-pages pages.json]

输出文档结构：
    封面页（标题 + 项目名 + 日期）→ 目录页（含页码，可两面）→ 正文

排版约定：
    正文 宋体/Times New Roman 小四（12pt）、1.5 倍行距、首行缩进 2 字符、两端对齐
    标题 黑体，H1 22pt、H2 16pt、H3 14pt、H4 12pt，使用 Word 内置 Heading 样式（可用导航窗格/书签）
    表格 Table Grid，表头加粗+浅蓝底纹、跨页重复、行内不拆页，9pt
    代码块 Consolas 8.5pt + 浅灰底纹（Mermaid 图保留源码）
    页脚居中页码；A4 页面、2.4cm 页边距

依赖：python-docx
"""

import datetime
import json
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CJK = re.compile(r'[\u4e00-\u9fff\u3000-\u303f\uff00-\uffef]')
HEADING = re.compile(r'^(#{1,6})\s+(.*?)\s*$')
FENCE = re.compile(r'^```\s*([A-Za-z0-9_+-]*)\s*$')
LIST_ITEM = re.compile(r'^(\s*)([-*+]|\d+[.)])\s+(.*)$')
TABLE_ROW = re.compile(r'^\s*\|.*$')
SEPARATOR_CELL = re.compile(r'^:?-{2,}:?$')
INLINE_TOKEN = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)\s]+\)|\*[^*\s][^*]*\*)')

BODY_FONT_EA = '宋体'
BODY_FONT_ASCII = 'Times New Roman'
HEAD_FONT_EA = '黑体'
MONO_FONT = 'Consolas'
HEADING_SIZES = {1: 22, 2: 16, 3: 14, 4: 12, 5: 12, 6: 11}
PROJECT_NAME = '智能求职辅助系统 · 课程综合实践'
CODE_FILL = 'F4F5F7'


# --------------------------------------------------------------------------- #
# 基础工具
# --------------------------------------------------------------------------- #
def plain(text):
    """去掉 Markdown 行内标记，得到纯文本（目录与书签匹配用）。"""
    text = re.sub(r'\[([^\]]+)\]\([^)\s]+\)', r'\1', text)
    text = text.replace('**', '').replace('`', '')
    text = re.sub(r'(?<!\*)\*([^*\s][^*]*)\*(?!\*)', r'\1', text)
    return text.strip()


def set_run_font(run, ea_font=BODY_FONT_EA, ascii_font=BODY_FONT_ASCII,
                 size=None, bold=None, italic=None, color=None):
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:ascii'), ascii_font)
    rfonts.set(qn('w:hAnsi'), ascii_font)
    rfonts.set(qn('w:eastAsia'), ea_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(element, fill):
    pr = element.get_or_add_pPr() if element.tag.endswith('}p') else element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pr.append(shd)


def add_bottom_border(paragraph, color='BFBFBF', size='4'):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size)
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pbdr.append(bottom)
    ppr.append(pbdr)


def thin_paragraph(doc, border=False, space_before=4, space_after=4):
    """生成一个很薄的空段落（用于水平分隔线/表格后的间隔，避免大片留白）。"""
    paragraph = doc.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Pt(0)
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = Pt(1)
    run = paragraph.add_run('')
    set_run_font(run, size=1)
    if border:
        add_bottom_border(paragraph)
    return paragraph


def add_page_number_footer(section):
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    run = paragraph.add_run()
    set_run_font(run, size=9, color=RGBColor(0x80, 0x80, 0x80))
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'PAGE'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def configure_document(doc, title):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    add_page_number_footer(section)

    normal = doc.styles['Normal']
    normal.font.name = BODY_FONT_ASCII
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), BODY_FONT_EA)
    fmt = normal.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(4)
    fmt.first_line_indent = Pt(24)

    core = doc.core_properties
    core.title = title
    core.author = '智能求职辅助系统课程小组'
    core.comments = '由 Markdown 自动转换生成'
    core.modified = datetime.datetime.now()


# --------------------------------------------------------------------------- #
# 行内格式
# --------------------------------------------------------------------------- #
def add_inline(paragraph, text, base_size=12, base_bold=False, base_color=None):
    pos = 0
    for match in INLINE_TOKEN.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=base_size, bold=base_bold, color=base_color)
        token = match.group(0)
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True, color=base_color)
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, ea_font=MONO_FONT, ascii_font=MONO_FONT,
                         size=max(base_size - 1.5, 8), bold=base_bold,
                         color=RGBColor(0xC7, 0x25, 0x4E))
        elif token.startswith('[') and '](' in token:
            label, url = token[1:-1].split('](', 1)
            run = paragraph.add_run(label)
            set_run_font(run, size=base_size, bold=base_bold, color=RGBColor(0x1F, 0x4E, 0x79))
            run.font.underline = True
            # 链接目标与链接文字重复时不再重复打印路径
            if os.path.basename(url) != label.strip() and url not in label:
                note = paragraph.add_run('（%s）' % url)
                set_run_font(note, size=max(base_size - 2, 8), color=RGBColor(0x88, 0x88, 0x88))
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=base_size, bold=base_bold, italic=True, color=base_color)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=base_size, bold=base_bold, color=base_color)


# --------------------------------------------------------------------------- #
# 块级解析
# --------------------------------------------------------------------------- #
def split_table_row(line):
    row = line.strip()
    if row.startswith('|'):
        row = row[1:]
    if row.endswith('|'):
        row = row[:-1]
    cells, buf, escaped = [], [], False
    for ch in row:
        if escaped:
            buf.append(ch)
            escaped = False
            continue
        if ch == '\\':
            escaped = True
            continue
        if ch == '|':
            cells.append(''.join(buf).strip())
            buf = []
            continue
        buf.append(ch)
    cells.append(''.join(buf).strip())
    return cells


def parse_blocks(text):
    lines = text.splitlines()
    blocks = []
    i, total = 0, len(lines)

    while i < total:
        line = lines[i]

        match = FENCE.match(line.strip())
        if match:
            i += 1
            code_lines = []
            while i < total and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1
            blocks.append(('code', match.group(1), code_lines))
            continue

        match = HEADING.match(line)
        if match:
            blocks.append(('heading', len(match.group(1)), match.group(2).strip()))
            i += 1
            continue

        if line.strip() in ('---', '***', '___'):
            blocks.append(('hr',))
            i += 1
            continue

        if TABLE_ROW.match(line):
            table_lines = []
            while i < total and TABLE_ROW.match(lines[i]):
                table_lines.append(lines[i])
                i += 1
            rows = [split_table_row(item) for item in table_lines]
            rows = [row for row in rows
                    if not all(SEPARATOR_CELL.match(cell) for cell in row if cell)]
            if rows:
                width = max(len(row) for row in rows)
                rows = [row + [''] * (width - len(row)) for row in rows]
                blocks.append(('table', rows))
            continue

        if LIST_ITEM.match(line):
            items = []
            while i < total:
                match = LIST_ITEM.match(lines[i])
                if not match:
                    break
                indent = len(match.group(1).expandtabs(2))
                marker = match.group(2)
                items.append((min(indent // 2, 3), marker[0].isdigit(), marker,
                              match.group(3).strip()))
                i += 1
            blocks.append(('list', items))
            continue

        if line.lstrip().startswith('>'):
            quote_lines = []
            while i < total and lines[i].lstrip().startswith('>'):
                quote_lines.append(re.sub(r'^\s*>\s?', '', lines[i]))
                i += 1
            blocks.append(('quote', quote_lines))
            continue

        if not line.strip():
            i += 1
            continue

        para_lines = [line.strip()]
        i += 1
        while i < total:
            nxt = lines[i]
            if (not nxt.strip() or HEADING.match(nxt) or FENCE.match(nxt.strip())
                    or TABLE_ROW.match(nxt) or LIST_ITEM.match(nxt)
                    or nxt.lstrip().startswith('>') or nxt.strip() in ('---', '***', '___')):
                break
            para_lines.append(nxt.strip())
            i += 1
        merged = ''
        for part in para_lines:
            if not merged:
                merged = part
            elif CJK.match(merged[-1]) or CJK.match(part[0]):
                merged += part
            else:
                merged += ' ' + part
        blocks.append(('para', merged))

    return blocks


# --------------------------------------------------------------------------- #
# 渲染
# --------------------------------------------------------------------------- #
def render_cover(doc, title_text):
    for _ in range(6):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        set_run_font(paragraph.add_run(''), size=12)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_after = Pt(16)
    run = paragraph.add_run(title_text)
    set_run_font(run, ea_font=HEAD_FONT_EA, size=26, bold=True)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(PROJECT_NAME)
    set_run_font(run, ea_font=HEAD_FONT_EA, size=14, color=RGBColor(0x44, 0x4A, 0x5A))

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    run = paragraph.add_run(datetime.date.today().strftime('%Y 年 %m 月 %d 日'))
    set_run_font(run, size=12, color=RGBColor(0x66, 0x6B, 0x7A))
    doc.add_page_break()


def render_toc(doc, entries, pages=None):
    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.first_line_indent = Pt(0)
    label.paragraph_format.space_before = Pt(0)
    label.paragraph_format.space_after = Pt(16)
    run = label.add_run('目  录')
    set_run_font(run, ea_font=HEAD_FONT_EA, size=18, bold=True)

    pages = pages or []
    for index, (level, text) in enumerate(entries):
        paragraph = doc.add_paragraph()
        fmt = paragraph.paragraph_format
        fmt.first_line_indent = Pt(0)
        fmt.left_indent = Cm(0.7 * (level - 2) if level > 2 else 0)
        fmt.space_after = Pt(4)
        fmt.line_spacing = 1.3
        fmt.tab_stops.add_tab_stop(Cm(16.2), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        run = paragraph.add_run(text)
        set_run_font(run, size=11.5 if level == 2 else 11, bold=(level == 2))
        page = None
        if isinstance(pages, dict):
            page = pages.get(text)
        elif isinstance(pages, list) and index < len(pages):
            page = pages[index]
        if page:
            tab_run = paragraph.add_run('\t')
            set_run_font(tab_run, size=11)
            number_run = paragraph.add_run(str(page))
            set_run_font(number_run, size=11)
    doc.add_page_break()


def render_heading(doc, level, text):
    paragraph = doc.add_paragraph(style='Heading %d' % level)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Pt(0)
    fmt.space_before = Pt(14 if level <= 2 else 8)
    fmt.space_after = Pt(8 if level <= 2 else 4)
    fmt.line_spacing = 1.2
    fmt.keep_with_next = True
    size = HEADING_SIZES.get(level, 11)
    add_inline(paragraph, text, base_size=size, base_bold=True)
    for run in paragraph.runs:
        set_run_font(run, ea_font=HEAD_FONT_EA, size=size, bold=True, color=RGBColor(0, 0, 0))


def render_paragraph(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(paragraph, text)


def render_code(doc, lines):
    paragraph = doc.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Pt(0)
    fmt.left_indent = Cm(0.4)
    fmt.right_indent = Cm(0.2)
    fmt.space_before = Pt(3)
    fmt.space_after = Pt(6)
    fmt.line_spacing = 1.05
    shade(paragraph._p, CODE_FILL)
    for index, code_line in enumerate(lines):
        if index:
            run = paragraph.add_run()
            run.add_break()
            set_run_font(run, ea_font=MONO_FONT, ascii_font=MONO_FONT, size=8.5)
        run = paragraph.add_run(code_line.replace('\t', '    '))
        set_run_font(run, ea_font=MONO_FONT, ascii_font=MONO_FONT, size=8.5,
                     color=RGBColor(0x1F, 0x29, 0x37))


def render_list(doc, items):
    for level, ordered, marker, text in items:
        paragraph = doc.add_paragraph()
        fmt = paragraph.paragraph_format
        fmt.left_indent = Cm(0.75 * (level + 1))
        fmt.first_line_indent = Cm(-0.45)
        fmt.space_after = Pt(2)
        fmt.line_spacing = 1.4
        prefix = (marker + ' ') if ordered else ('• ' if level == 0 else '◦ ')
        run = paragraph.add_run(prefix)
        set_run_font(run, size=12)
        add_inline(paragraph, text)


def render_quote(doc, lines):
    text = ''
    for part in lines:
        if not text:
            text = part
        elif CJK.match(text[-1]) or (part and CJK.match(part[0])):
            text += part
        else:
            text += ' ' + part
    paragraph = doc.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Pt(0)
    fmt.left_indent = Cm(0.6)
    fmt.space_before = Pt(4)
    fmt.space_after = Pt(6)
    shade(paragraph._p, 'EEF2FF')
    add_inline(paragraph, text, base_size=11, base_color=RGBColor(0x37, 0x41, 0x51))


def render_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    head_cells = table.rows[0].cells
    for index, text in enumerate(header):
        cell = head_cells[index]
        paragraph = cell.paragraphs[0]
        fmt = paragraph.paragraph_format
        fmt.first_line_indent = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = 1.15
        add_inline(paragraph, text, base_size=9, base_bold=True)
        for run in paragraph.runs:
            if MONO_FONT not in (run.font.name or ''):
                set_run_font(run, ea_font=HEAD_FONT_EA, size=9, bold=True)
        shade(cell._tc, 'DCE4F5')

    for row in body:
        cells = table.add_row().cells
        for index, text in enumerate(row):
            paragraph = cells[index].paragraphs[0]
            fmt = paragraph.paragraph_format
            fmt.first_line_indent = Pt(0)
            fmt.space_after = Pt(0)
            fmt.line_spacing = 1.15
            add_inline(paragraph, text, base_size=9)

    for row_index, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit')
        trpr.append(cant)
        if row_index == 0:
            tbl_header = OxmlElement('w:tblHeader')
            tbl_header.set(qn('w:val'), 'true')
            trpr.append(tbl_header)

    thin_paragraph(doc, space_before=2, space_after=2)


# --------------------------------------------------------------------------- #
# 组装文档
# --------------------------------------------------------------------------- #
def build_document(text, doc_title, toc_pages=None):
    doc = Document()
    configure_document(doc, doc_title)

    blocks = parse_blocks(text)
    title_index = None
    for index, block in enumerate(blocks):
        if block[0] == 'heading' and block[1] == 1:
            title_index = index
            break

    if title_index is not None:
        title_text = plain(blocks[title_index][2])
        body_blocks = blocks[:title_index] + blocks[title_index + 1:]
    else:
        title_text = doc_title
        body_blocks = blocks

    render_cover(doc, title_text)

    toc_entries = [(block[1], plain(block[2])) for block in body_blocks
                   if block[0] == 'heading' and 2 <= block[1] <= 3]
    render_toc(doc, toc_entries, toc_pages)

    for block in body_blocks:
        kind = block[0]
        if kind == 'heading':
            render_heading(doc, block[1], block[2])
        elif kind == 'para':
            render_paragraph(doc, block[1])
        elif kind == 'code':
            render_code(doc, block[2])
        elif kind == 'table':
            render_table(doc, block[1])
        elif kind == 'list':
            render_list(doc, block[1])
        elif kind == 'quote':
            render_quote(doc, block[1])
        elif kind == 'hr':
            thin_paragraph(doc, border=True, space_before=6, space_after=6)

    return doc


def convert(src_path, dst_path, toc_pages=None):
    with open(src_path, 'r', encoding='utf-8') as handle:
        text = handle.read()
    title = os.path.splitext(os.path.basename(src_path))[0]
    doc = build_document(text, title, toc_pages)
    doc.save(dst_path)
    return doc


def main():
    args = [item for item in sys.argv[1:] if not item.startswith('--')]
    if len(args) < 2:
        print('用法：python md2docx.py <源目录> <输出目录> [--toc-pages pages.json]')
        return 2
    src_dir, out_dir = args[0], args[1]
    pages_map = None
    for index, item in enumerate(sys.argv):
        if item == '--toc-pages' and index + 1 < len(sys.argv):
            with open(sys.argv[index + 1], 'r', encoding='utf-8') as handle:
                pages_map = json.load(handle)

    os.makedirs(out_dir, exist_ok=True)
    files = sorted(name for name in os.listdir(src_dir) if name.lower().endswith('.md'))
    for name in files:
        src_path = os.path.join(src_dir, name)
        dst_path = os.path.join(out_dir, os.path.splitext(name)[0] + '.docx')
        pages = None
        if pages_map is not None:
            pages = pages_map.get(os.path.splitext(name)[0])
        doc = convert(src_path, dst_path, pages)
        print('%s -> %s（段落 %d，表格 %d）' % (
            name, os.path.basename(dst_path), len(doc.paragraphs), len(doc.tables)))
    print('完成，共 %d 个文件，输出目录：%s' % (len(files), out_dir))
    return 0


if __name__ == '__main__':
    sys.exit(main())
